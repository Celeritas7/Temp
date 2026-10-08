#!/usr/bin/env python3
"""Build the BRK-100_Tooling sample project used to try out DocSearch.

The project is a made-up motor bracket assembly (BRK-100) and the tools used to
assemble it. Every document refers to the same tool numbers, so one search finds
a tool across Word, Excel, PDF, PowerPoint, STEP and text files.

Only needed if you want to rebuild the samples (they are already in the repo):
    pip install cadquery python-docx openpyxl python-pptx reportlab
    python samples/make_samples.py
"""

import csv
import os
import sys

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BRK-100_Tooling")
sys.path.insert(0, os.path.join(ROOT, "CAD"))
import tool_models  # noqa: E402  (parametric CAD; also gives the drawings their dimensions)

FX, PT, GA = tool_models.FX201, tool_models.PT301, tool_models.GA401

TOOLS = [
    # id, name, kind, material / maker, CAD file, location, calibration due
    ("FX-201", "Locating Fixture", "Made in-house", "Aluminium 6061-T6 plate, steel dowel pins D8 h6",
     "CAD/FX-201_Locating_Fixture.step", "Cell 3, bench A", "n/a (visual check daily)"),
    ("PT-301", "Bushing Press Punch", "Made in-house", "Tool steel 1.2379, hardened 58-62 HRC",
     "CAD/PT-301_Bushing_Press_Punch.step", "Cell 3, press AP-501", "n/a (pilot D checked weekly)"),
    ("GA-401", "Go/No-Go Pin Gauge", "Made in-house", "Gauge steel, GO D10.00 / NO-GO D10.05",
     "CAD/GA-401_Go_No-Go_Pin_Gauge.step", "Cell 3, gauge board", "2027-03-01"),
    ("TW-101", "Torque Wrench 2-25 Nm", "Purchased", "Click-type, 1/4in square drive",
     "-", "Cell 3, tool board", "2026-12-15"),
    ("AP-501", "Arbor Press 1 t", "Purchased", "Manual rack-and-pinion press",
     "-", "Cell 3, bench A", "n/a"),
]

STEPS = [
    ("10", "Clean locating fixture FX-201 and check both dowel pins for burrs or damage.", "FX-201"),
    ("20", "Place bracket BRK-100 in the FX-201 nest so that both dowel pins enter the bracket holes. "
           "The bracket must sit flat on the pocket floor.", "FX-201"),
    ("30", "Put bushing BU-12 on the pilot of press punch PT-301. Press it into the bracket bore with "
           "arbor press AP-501 until the punch shoulder seats on the bushing flange. Maximum force 4 kN.",
     "PT-301, AP-501"),
    ("40", "Check the bushing bore with pin gauge GA-401: the GO end (D10.00) must enter freely, the NO-GO "
           "end (D10.05) must not enter. Record the result in the inspection log.", "GA-401"),
    ("50", "Fit motor plate MP-50. Apply medium-strength threadlocker and tighten the 4 x M6 screws SC-616 "
           "to 9.5 Nm in a cross pattern with torque wrench TW-101.", "TW-101"),
    ("60", "Lift the assembly out of FX-201 using the finger slots and attach the serial number label.",
     "FX-201"),
]


def write_docx(path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Work Instruction WI-BRK-100-A  Rev C", 0)
    doc.add_paragraph("Product: BRK-100 Motor Bracket Assembly    Cell: 3    Takt time: 4 min")
    doc.add_heading("Safety", 1)
    doc.add_paragraph("Wear safety glasses. Keep hands clear of the arbor press ram while pressing.")
    doc.add_paragraph("Do not use press punch PT-301 if the pilot is chipped or worn (see NCR-0042).")
    doc.add_heading("Tools required", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Tool ID", "Description", "Location")):
        cell.text = text
    for tool_id, name, _, _, _, location, _ in TOOLS:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = tool_id, name, location
    doc.add_heading("Assembly steps", 1)
    steps = doc.add_table(rows=1, cols=3)
    steps.style = "Light Grid Accent 1"
    for cell, text in zip(steps.rows[0].cells, ("Step", "Operation", "Tools")):
        cell.text = text
    for number, text, tools in STEPS:
        row = steps.add_row().cells
        row[0].text, row[1].text, row[2].text = number, text, tools
    doc.add_heading("Revision history", 1)
    doc.add_paragraph("Rev A  2025-11-03  First issue.")
    doc.add_paragraph("Rev B  2026-04-22  Torque changed from 8 Nm to 9.5 Nm (ECR-118).")
    doc.add_paragraph("Rev C  2026-09-21  Weekly pilot check of PT-301 added after NCR-0042.")
    doc.save(path)


def write_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tooling Register"
    header = ["Tool ID", "Description", "Type", "Material / Details", "CAD file", "Location", "Calibration due"]
    ws.append(header)
    for tool in TOOLS:
        ws.append(list(tool))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FDE")
    for column, width in zip("ABCDEFG", (10, 26, 15, 46, 38, 22, 28)):
        ws.column_dimensions[column].width = width

    cal = wb.create_sheet("Calibration Log")
    cal.append(["Date", "Tool ID", "Check", "Measured", "Limit", "Result", "By"])
    for row in [
        ("2025-12-15", "TW-101", "Torque at 9.5 Nm", "9.41 Nm", "9.5 Nm +/-4%", "PASS", "Cal lab"),
        ("2026-03-01", "GA-401", "GO pin diameter", "10.001 mm", "10.000 +0.002/-0", "PASS", "Cal lab"),
        ("2026-03-01", "GA-401", "NO-GO pin diameter", "10.049 mm", "10.050 +0/-0.002", "PASS", "Cal lab"),
        ("2026-09-14", "PT-301", "Pilot diameter", "9.96 mm", "10.02 +/-0.01", "FAIL - replaced", "Toolroom"),
        ("2026-09-15", "PT-301", "Pilot diameter (new punch)", "10.021 mm", "10.02 +/-0.01", "PASS", "Toolroom"),
    ]:
        cal.append(row)
    for cell in cal[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFG", (12, 10, 28, 12, 20, 16, 10)):
        cal.column_dimensions[column].width = width
    wb.save(path)


def write_pptx(path):
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "BRK-100 Operator Training"
    title.placeholders[1].text = "Assembly tools and how to use them - Cell 3"

    def bullets(heading, lines, notes=None):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        body.text = lines[0]
        for line in lines[1:]:
            body.add_paragraph().text = line
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    bullets("FX-201 Locating Fixture", [
        f"Holds bracket BRK-100 on two D{FX['dowel_diameter']:g} dowel pins",
        "Check the pins every shift - a bent pin gives a crooked bushing",
        "Use the finger slots to lift the part out, never a screwdriver",
    ], notes="Show the operator where the spare dowel pins are kept.")
    bullets("PT-301 Bushing Press Punch", [
        f"Pilot D{PT['pilot_diameter']:g} guides bushing BU-12 straight into the bore",
        "Press with AP-501 until the shoulder seats - do not exceed 4 kN",
        "Worn pilot = loose bushings (NCR-0042). Pilot is checked weekly",
    ])
    bullets("GA-401 Go/No-Go Pin Gauge", [
        f"GO end D{GA['go_diameter']:.2f} must enter the bore freely",
        f"NO-GO end D{GA['nogo_diameter']:.2f} (groove on handle) must NOT enter",
        "Record every check in the inspection log",
    ])
    bullets("TW-101 Torque Wrench", [
        "Set to 9.5 Nm for the M6 screws SC-616",
        "Tighten in a cross pattern, stop at the first click",
        "Return the setting to 2 Nm after the shift to protect the spring",
    ])
    bullets("Common mistakes", [
        "Pressing the bushing without the part seated on both dowel pins",
        "Using the NO-GO end first and forcing it into the bore",
        "Torque wrench TW-101 left at a high setting overnight",
    ])
    prs.save(path)


def write_pdf(path):
    """Simple dimensioned tool drawings, one tool per page, generated from the CAD parameters."""
    c = canvas.Canvas(path, pagesize=landscape(A4))
    page_w, page_h = landscape(A4)
    scale = 2.0  # drawing units per mm (roughly 1:1.4 on A4)

    def frame(tool_id, title, material, cad_file, scale_text):
        c.setLineWidth(1.2)
        c.rect(20, 20, page_w - 40, page_h - 40)
        c.setLineWidth(0.8)
        x0, y0, w, h = page_w - 330, 20, 310, 80
        c.rect(x0, y0, w, h)
        c.line(x0, y0 + 50, x0 + w, y0 + 50)
        c.line(x0 + 200, y0, x0 + 200, y0 + 50)
        c.setFont("Helvetica-Bold", 14)
        c.drawString(x0 + 8, y0 + 58, f"{tool_id}  {title}")
        c.setFont("Helvetica", 8)
        c.drawString(x0 + 8, y0 + 34, f"Material: {material}")
        c.drawString(x0 + 8, y0 + 22, "Used on: BRK-100 Motor Bracket Assembly")
        c.drawString(x0 + 8, y0 + 10, "All dimensions in mm. General tolerance ISO 2768-m")
        c.drawString(x0 + 208, y0 + 34, f"Drawing: DRW-{tool_id}")
        c.drawString(x0 + 208, y0 + 22, f"Rev: C   Scale: {scale_text}")
        c.setFont("Helvetica", 6.5)
        c.drawString(x0 + 8, y0 + 42, f"CAD model: {cad_file}")

    def hdim(x1, x2, y, text):
        c.setLineWidth(0.4)
        c.line(x1, y, x2, y)
        for x, d in ((x1, 1), (x2, -1)):
            c.line(x, y, x + 5 * d, y + 2)
            c.line(x, y, x + 5 * d, y - 2)
        c.setFont("Helvetica", 8)
        c.drawCentredString((x1 + x2) / 2, y + 3, text)

    def vdim(x, y1, y2, text):
        c.setLineWidth(0.4)
        c.line(x, y1, x, y2)
        for y, d in ((y1, 1), (y2, -1)):
            c.line(x, y, x - 2, y + 5 * d)
            c.line(x, y, x + 2, y + 5 * d)
        c.saveState()
        c.translate(x - 3, (y1 + y2) / 2)
        c.rotate(90)
        c.setFont("Helvetica", 8)
        c.drawCentredString(0, 0, text)
        c.restoreState()

    # ---- page 1: FX-201 top view
    frame("FX-201", "Locating Fixture", "Aluminium 6061-T6 / dowels steel h6", TOOLS[0][4], "1:1.4")
    cx, cy = 300, 330
    L, W = FX["plate_length"] * scale, FX["plate_width"] * scale
    c.setLineWidth(1)
    c.roundRect(cx - L / 2, cy - W / 2, L, W, FX["corner_radius"] * scale)
    pl, pw = FX["pocket_length"] * scale, FX["pocket_width"] * scale
    c.roundRect(cx - pl / 2, cy - pw / 2, pl, pw, FX["pocket_corner_radius"] * scale)
    sw, sl = FX["relief_slot_width"] * scale, FX["relief_slot_length"] * scale
    for d in (-1, 1):  # finger slots
        y_edge = cy + d * pw / 2
        c.rect(cx - sw / 2, min(y_edge, y_edge + d * sl), sw, sl)
    for d in (-1, 1):
        c.circle(cx + d * FX["dowel_spacing"] / 2 * scale, cy, FX["dowel_diameter"] / 2 * scale)
    dx, dy = (FX["plate_length"] / 2 - FX["mount_hole_inset"]) * scale, (FX["plate_width"] / 2 - FX["mount_hole_inset"]) * scale
    for sx in (-1, 1):
        for sy in (-1, 1):
            c.circle(cx + sx * dx, cy + sy * dy, FX["mount_hole_diameter"] / 2 * scale)
            c.circle(cx + sx * dx, cy + sy * dy, FX["mount_cbore_diameter"] / 2 * scale)
    hdim(cx - L / 2, cx + L / 2, cy + W / 2 + 20, f"{FX['plate_length']:g}")
    vdim(cx - L / 2 - 20, cy - W / 2, cy + W / 2, f"{FX['plate_width']:g}")
    hdim(cx - pl / 2, cx + pl / 2, cy - W / 2 - 18, f"Pocket {FX['pocket_length']:g} x {FX['pocket_width']:g} deep {FX['pocket_depth']:g}")
    hdim(cx - FX["dowel_spacing"] / 2 * scale, cx + FX["dowel_spacing"] / 2 * scale, cy + 14, f"{FX['dowel_spacing']:g}")
    c.setFont("Helvetica", 9)
    notes = [
        "NOTES",
        f"1. 2x dowel pin D{FX['dowel_diameter']:g} h6, pressed, {FX['dowel_height']:g} above pocket floor.",
        f"2. 4x M6 counterbore D{FX['mount_cbore_diameter']:g} x {FX['mount_cbore_depth']:g} deep, thru D{FX['mount_hole_diameter']:g}.",
        f"3. Plate thickness {FX['plate_thickness']:g}. Break all edges {FX['edge_chamfer']:g} x 45 deg.",
        f"4. Finger slots {FX['relief_slot_width']:g} wide both sides of pocket.",
        "5. Nest fits bracket BRK-100 (60 x 40) with 1 mm clearance per side.",
    ]
    for i, line in enumerate(notes):
        c.drawString(530, 480 - i * 14, line)
    c.showPage()

    # ---- page 2: PT-301 side view
    frame("PT-301", "Bushing Press Punch", "Tool steel 1.2379, 58-62 HRC", TOOLS[1][4], "1.4:1")
    scale = 4.0
    x, y = 150, 330
    segs = [(PT["pilot_length"], PT["pilot_diameter"], "Pilot"), (PT["shoulder_length"], PT["shoulder_diameter"], "Shoulder"),
            (PT["shank_length"], PT["shank_diameter"], "Shank")]
    for length, dia, label in segs:
        c.setLineWidth(1)
        c.rect(x, y - dia / 2 * scale, length * scale, dia * scale)
        hdim(x, x + length * scale, y - PT["shoulder_diameter"] / 2 * scale - 25, f"{length:g}")
        c.setFont("Helvetica-Bold", 9)
        c.drawCentredString(x + length * scale / 2, y + PT["shoulder_diameter"] / 2 * scale + 14 + (12 if label == "Pilot" else 0),
                            f"{label} D{dia:g}")
        x += length * scale
    c.setDash(6, 3)
    c.line(130, y, x + 20, y)
    c.setDash()
    c.setFont("Helvetica", 9)
    for i, line in enumerate([
        "NOTES",
        f"1. Pilot D{PT['pilot_diameter']:g} +/-0.01 - check weekly, replace below D10.01 (NCR-0042).",
        f"2. Chamfer pilot {PT['pilot_chamfer']:g} x 45 deg, shank {PT['shank_chamfer']:g} x 45 deg.",
        f"3. 2 flats on shank, {PT['grip_flat_depth']:g} deep x {PT['grip_flat_length']:g} long.",
        "4. Use with arbor press AP-501. Max force 4 kN.",
    ]):
        c.drawString(150, 170 - i * 14, line)
    c.showPage()

    # ---- page 3: GA-401 side view
    frame("GA-401", "Go/No-Go Pin Gauge", "Gauge steel, hardened", TOOLS[2][4], "1.4:1")
    scale = 4.0
    x, y = 80, 330
    segs = [(GA["go_length"], GA["go_diameter"], "GO"), (GA["neck_length"], GA["neck_diameter"], ""),
            (GA["handle_length"], GA["handle_diameter"], "Handle"), (GA["neck_length"], GA["neck_diameter"], ""),
            (GA["nogo_length"], GA["nogo_diameter"], "NO-GO")]
    for length, dia, label in segs:
        c.setLineWidth(1)
        c.rect(x, y - dia / 2 * scale, length * scale, dia * scale)
        if label:
            hdim(x, x + length * scale, y - GA["handle_diameter"] / 2 * scale - 25, f"{length:g}")
            c.setFont("Helvetica-Bold", 9)
            c.drawCentredString(x + length * scale / 2, y + GA["handle_diameter"] / 2 * scale + 12,
                                f"{label} D{dia:.2f}" if label != "Handle" else f"Handle D{dia:g}")
        x += length * scale
    c.setFont("Helvetica", 9)
    for i, line in enumerate([
        "NOTES",
        f"1. GO D{GA['go_diameter']:.3f} +0.002/-0, NO-GO D{GA['nogo_diameter']:.3f} +0/-0.002.",
        f"2. Identification groove {GA['groove_width']:g} wide x {GA['groove_depth']:g} deep next to the NO-GO end.",
        "3. Checks the BU-12 bushing bore D10.00 +0.05 after pressing (step 40).",
        "4. Calibration every 12 months - see Tooling Register.",
    ]):
        c.drawString(120, 200 - i * 14, line)
    c.showPage()
    c.save()


def write_text_files():
    with open(os.path.join(ROOT, "Quality", "NCR-0042_Loose_Bushing.txt"), "w", encoding="utf-8") as fh:
        fh.write("""NONCONFORMANCE REPORT  NCR-0042
Date raised:  2026-09-14
Product:      BRK-100 Motor Bracket Assembly
Found at:     Final inspection, 3 of 120 units

Description
Bushing BU-12 could be turned by hand in the bracket bore on 3 units (serials 26-0871, 26-0874, 26-0880).
The bushing bore passed the GA-401 go/no-go check on all 3 units.

Root cause
The pilot of press punch PT-301 was worn to 9.96 mm (drawing 10.02 +/-0.01). The bushing tilted while
being pressed and the press fit was damaged on one side.

Containment
All units built since the last good PT-301 check (2026-09-07) re-checked: 3 rejects found, scrapped.

Corrective action
1. PT-301 replaced with a new punch (pilot measured 10.021 mm on 2026-09-15).
2. Weekly pilot diameter check of PT-301 added to WI-BRK-100-A Rev C.
3. Spare PT-301 punch kept in the toolroom.

Status: CLOSED 2026-09-21
""")
    with open(os.path.join(ROOT, "Quality", "Inspection_Log_2026-09.csv"), "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Serial", "Date", "GA-401 GO", "GA-401 NO-GO", "Torque 9.5 Nm (TW-101)", "Result", "Operator"])
        rows = [
            ("26-0868", "2026-09-12", "enters", "does not enter", "OK", "PASS", "AM"),
            ("26-0871", "2026-09-12", "enters", "does not enter", "OK", "FAIL - loose bushing, NCR-0042", "AM"),
            ("26-0874", "2026-09-13", "enters", "does not enter", "OK", "FAIL - loose bushing, NCR-0042", "RK"),
            ("26-0880", "2026-09-13", "enters", "does not enter", "OK", "FAIL - loose bushing, NCR-0042", "RK"),
            ("26-0881", "2026-09-16", "enters", "does not enter", "OK", "PASS", "AM"),
            ("26-0882", "2026-09-16", "enters", "enters", "OK", "FAIL - bore oversize", "AM"),
            ("26-0883", "2026-09-16", "enters", "does not enter", "OK", "PASS", "AM"),
        ]
        writer.writerows(rows)
    with open(os.path.join(ROOT, "Meetings", "2026-09-20_Tooling_Review.md"), "w", encoding="utf-8") as fh:
        fh.write("""# Tooling review - Cell 3 - 2026-09-20

Attendees: production, quality, toolroom

## NCR-0042 follow-up
- New PT-301 punch is in use since 2026-09-15. No loose bushings since then.
- Toolroom to make one spare PT-301 and add the pilot check to the weekly tool check list.

## BRK-200 variant (longer bracket, 80 x 40)
- Needs a new locating fixture FX-202. Start from the FX-201 CAD model: change pocket_length
  from 62 to 82 and dowel_spacing from 44 to 64 in tool_models.py.
- PT-301 and GA-401 can be shared between BRK-100 and BRK-200 (same bushing BU-12).

## Torque wrench
- TW-101 calibration due 2026-12-15. Book the cal lab in November.
- Proposal: second torque wrench TW-102 for the BRK-200 line.

## Actions
| Action | Owner | Due |
|---|---|---|
| Make spare PT-301 | Toolroom | 2026-10-05 |
| Design FX-202 fixture for BRK-200 | Tooling engineer | 2026-10-15 |
| Book TW-101 calibration | Quality | 2026-11-15 |
""")


def main():
    for folder in ("CAD", "Drawings", "Work_Instructions", "Tooling", "Training", "Quality", "Meetings"):
        os.makedirs(os.path.join(ROOT, folder), exist_ok=True)
    tool_models.main(os.path.join(ROOT, "CAD"))
    write_docx(os.path.join(ROOT, "Work_Instructions", "WI-BRK-100-A_Assembly_RevC.docx"))
    write_xlsx(os.path.join(ROOT, "Tooling", "Tooling_Register.xlsx"))
    write_pptx(os.path.join(ROOT, "Training", "BRK-100_Operator_Training.pptx"))
    write_pdf(os.path.join(ROOT, "Drawings", "BRK-100_Tool_Drawings.pdf"))
    write_text_files()
    print("Sample project written to", ROOT)


if __name__ == "__main__":
    main()
